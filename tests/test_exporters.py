import io

from docx import Document

from handnotes.exporters import to_docx, to_pdf


def test_to_pdf_returns_pdf_bytes_even_with_curly_quotes():
    data = to_pdf("Notes", "Denise’s “avant-garde” writing — really", "Short summary")
    assert data.startswith(b"%PDF")


def test_to_docx_contains_title_text_and_summary():
    data = to_docx("Notes", "line one\nline two", "the summary")
    paragraphs = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Notes" in paragraphs
    assert "line one" in paragraphs
    assert "the summary" in paragraphs


def test_to_docx_skips_summary_section_when_empty():
    data = to_docx("Notes", "text", "")
    paragraphs = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Summary" not in paragraphs
