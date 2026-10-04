"""Turn the corrected notes (and optional summary) into PDF or DOCX bytes."""

import io

from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos

# fpdf's built-in fonts are latin-1 only, so map common typographic characters.
_LATIN1_REPLACEMENTS = {
    "’": "'", "‘": "'", "“": '"', "”": '"', "—": "-", "–": "-",
    "…": "...", "→": "->", "•": "-",
}


def _plain(text: str) -> str:
    """Drop markdown bold markers the model likes to add."""
    return text.replace("**", "")


def _latin1(text: str) -> str:
    for original, replacement in _LATIN1_REPLACEMENTS.items():
        text = text.replace(original, replacement)
    return text.encode("latin-1", "replace").decode("latin-1")


def _pdf_section(pdf: FPDF, heading: str, body: str) -> None:
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, heading, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, _latin1(_plain(body)), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)


def to_pdf(title: str, text: str, summary: str = "") -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _latin1(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    _pdf_section(pdf, "Transcription", text)
    if summary:
        _pdf_section(pdf, "Summary", summary)
    return bytes(pdf.output())


def to_docx(title: str, text: str, summary: str = "") -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    document.add_heading("Transcription", level=1)
    for line in text.splitlines():
        document.add_paragraph(line)
    if summary:
        document.add_heading("Summary", level=1)
        for line in _plain(summary).splitlines():
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
